"""Note reading, autosave, inline AI actions and export (PRD §7.2)."""

from __future__ import annotations

import functools
import io
import logging
from typing import Annotated, Any, Literal

from fastapi import APIRouter, Header, Query
from fastapi.responses import Response
from pydantic import BaseModel, Field
from sqlalchemy import select

from .. import prompts
from ..config import get_config
from ..db.enums import JobStatus, JobType
from ..db.models import Job, Note, NoteBlock, Segment, Source
from ..db.session import session_scope
from ..errors import ConflictError, NotFoundError, ValidationError
from ..pipeline import tiptap
from ..pipeline.generate.notes import SIMPLIFIED_STYLE, reindex_note
from ..providers import gateway
from ..providers.base import text_message
from .common import note_dict, run_db, segment_dict

log = logging.getLogger(__name__)
router = APIRouter(prefix="/notes", tags=["notes"])

AiAction = Literal["explain", "expand", "simplify", "key_points", "rewrite"]


class NotePatch(BaseModel):
    doc: dict[str, Any] | None = None
    title: str | None = None
    #: Local patch: the note version the editor loaded. Regeneration bumps it, so a save from a
    #: tab still showing the previous generation is refused instead of overwriting the new one.
    base_version: int | None = None


def _top_level_texts(doc: dict[str, Any]) -> list[tuple[str, str]]:
    return [
        (str(node.get("type")), tiptap._collect_text(node).strip())
        for node in doc.get("content") or []
    ]


def _blanked_blocks(old: dict[str, Any], new: dict[str, Any]) -> int:
    """Blocks that kept their type and position but lost all their text.

    Local patch: the signature of the editor silently dropping node content (empty callouts)
    and saving it. A person deletes text one block at a time, and removing whole blocks shifts
    the structure, so two or more blanked in one debounced save is not an edit.
    """
    before, after = _top_level_texts(old), _top_level_texts(new)
    if len(before) != len(after):
        return 0
    return sum(
        1
        for (old_type, old_text), (new_type, new_text) in zip(before, after, strict=True)
        if old_type == new_type and old_text and not new_text
    )


class AiActionRequest(BaseModel):
    action: AiAction = "explain"
    selection: str = Field(min_length=1, max_length=8000)
    context: str = ""


@router.get("/{note_id}")
def get_note(note_id: int) -> dict[str, Any]:
    with session_scope() as session:
        note = session.get(Note, note_id)
        if note is None:
            raise NotFoundError(f"Note {note_id} not found.")
        blocks = list(
            session.execute(
                select(NoteBlock).where(NoteBlock.note_id == note_id).order_by(NoteBlock.index)
            ).scalars()
        )
        payload = note_dict(note)
        payload["blocks"] = [
            {
                "id": block.id,
                "index": block.index,
                "type": str(block.type),
                "text": block.text,
                "segment_ids": list(block.segment_ids or []),
            }
            for block in blocks
        ]
        return payload


@functools.cache
def _served_build() -> str | None:
    """The UI build being served, or None in development where Vite serves it."""
    path = get_config().static_dir / "build-id.txt"
    return path.read_text().strip() if path.exists() else None


def _guard_doc_write(
    session: Any, note: Note, payload: NotePatch, client_build: str | None
) -> None:
    """Local patch: refuse autosaves that can only be accidents (see LOCAL_SETUP.md)."""
    served = _served_build()
    if served is not None and client_build != served:
        # A tab opened before a deploy keeps running the old code; its saves are what emptied
        # callouts after the fix had shipped.
        log.warning("refused save of note %s from stale UI build %r", note.id, client_build)
        raise ConflictError("QuickLearn was updated since this page loaded. Reload the page.")
    # Only jobs that rewrite this note lock it: the simplified note generating must not stop
    # edits to the standard one.
    writers = [JobType.INGEST]
    writers.append(
        JobType.GENERATE_SIMPLIFIED_NOTES
        if note.style == SIMPLIFIED_STYLE
        else JobType.GENERATE_NOTES
    )
    busy = session.execute(
        select(Job.id).where(
            Job.source_id == note.source_id,
            Job.type.in_(writers),
            Job.status.in_([JobStatus.QUEUED, JobStatus.RUNNING]),
        )
    ).first()
    if busy is not None:
        raise ConflictError("This source is still processing — edits are paused until it finishes.")
    if payload.base_version is not None and payload.base_version != note.version:
        raise ConflictError("This note was regenerated since you opened it. Reload to see it.")
    blanked = _blanked_blocks(note.doc or {}, payload.doc or {})
    if blanked >= 2:
        log.warning("refused save of note %s: %d blocks lost all their text", note.id, blanked)
        raise ConflictError("Save refused: it would have erased the text of several blocks.")


@router.patch("/{note_id}")
async def patch_note(
    note_id: int,
    payload: NotePatch,
    x_quicklearn_build: Annotated[str | None, Header()] = None,
) -> dict[str, Any]:
    """Autosave. The client debounces at 800ms; this just has to be quick and idempotent."""

    def save() -> dict[str, Any]:
        with session_scope() as session:
            note = session.get(Note, note_id)
            if note is None:
                raise NotFoundError(f"Note {note_id} not found.")
            if payload.doc is not None:
                if payload.doc.get("type") != "doc":
                    raise ValidationError("Document must be a TipTap doc node.")
                _guard_doc_write(session, note, payload, x_quicklearn_build)
                note.doc = payload.doc
                # Marks the note as hand-edited, so regeneration warns before overwriting
                # rather than silently discarding the user's work (PRD §12.3).
                note.is_dirty = True
            if payload.title is not None:
                note.title = payload.title.strip()[:400] or note.title
            return note_dict(note, include_doc=False)

    result = await run_db(save)
    if payload.doc is not None:
        # Keep search and citation anchors in step with the edit.
        await run_db(reindex_note, note_id)
    return result


@router.get("/{note_id}/segments")
def note_segments(note_id: int) -> dict[str, Any]:
    """Segments backing this note, for resolving gutter anchors."""
    with session_scope() as session:
        note = session.get(Note, note_id)
        if note is None:
            raise NotFoundError(f"Note {note_id} not found.")
        rows = list(
            session.execute(
                select(Segment).where(Segment.source_id == note.source_id).order_by(Segment.index)
            ).scalars()
        )
    return {"segments": [segment_dict(segment) for segment in rows]}


@router.post("/{note_id}/ai-action")
async def ai_action(note_id: int, payload: AiActionRequest) -> dict[str, Any]:
    """Inline AI action on an editor selection."""

    def load() -> int | None:
        with session_scope() as session:
            note = session.get(Note, note_id)
            if note is None:
                raise NotFoundError(f"Note {note_id} not found.")
            return note.source_id

    source_id = await run_db(load)
    prompt = prompts.render(
        "ai_action",
        action=payload.action,
        selection=payload.selection,
        context=payload.context[:4000],
    )
    text = await gateway.complete_text(
        f"ai_action:{payload.action}",
        [text_message("user", prompt)],
        temperature=0.4,
        max_tokens=1500,
        source_id=source_id,
    )
    return {"action": payload.action, "result": text.strip()}


@router.post("/{note_id}/make-cards")
async def make_cards_from_selection(note_id: int, payload: AiActionRequest) -> dict[str, Any]:
    """Turn a selection into flashcards without leaving the editor."""
    from ..db.enums import CardType
    from ..db.models import Card
    from ..db.types import utcnow
    from ..pipeline.schemas import CardList

    def load() -> tuple[int, str]:
        with session_scope() as session:
            note = session.get(Note, note_id)
            if note is None:
                raise NotFoundError(f"Note {note_id} not found.")
            source = session.get(Source, note.source_id)
            return note.source_id, (source.title if source else "")

    source_id, title = await run_db(load)
    prompt = prompts.render(
        "cards",
        title=title,
        concepts=[
            {
                "id": 0,
                "name": "Selected passage",
                "definition": payload.selection,
                "importance": 4,
                "segment_ids": [],
            }
        ],
        max_per_concept=4,
    )
    result = await gateway.complete_structured(
        "cards from selection",
        [text_message("user", prompt)],
        CardList,
        temperature=0.3,
        source_id=source_id,
    )

    def persist() -> list[int]:
        now = utcnow()
        with session_scope() as session:
            rows = [
                Card(
                    source_id=source_id,
                    type=CardType(card.type),
                    front=card.front.strip(),
                    back=card.back.strip(),
                    due=now,
                )
                for card in result.cards
                if card.front.strip() and card.back.strip()
            ]
            session.add_all(rows)
            session.flush()
            return [row.id for row in rows]

    ids = await run_db(persist)
    return {"created": len(ids), "card_ids": ids}


# --------------------------------------------------------------------------- export


@router.get("/{note_id}/export")
def export_note(
    note_id: int,
    format: Annotated[Literal["md", "txt", "docx", "html", "json"], Query()] = "md",
) -> Response:
    """Export a note.

    PDF is deliberately absent: it is produced by the browser's own print-to-PDF from the
    `?print` view, rather than by shipping a headless Chromium in the image for one format.
    """
    with session_scope() as session:
        note = session.get(Note, note_id)
        if note is None:
            raise NotFoundError(f"Note {note_id} not found.")
        doc = note.doc or tiptap.empty_doc()
        title = note.title

    safe_title = "".join(c for c in title if c.isalnum() or c in " -_").strip() or "note"

    if format == "md":
        body = tiptap.doc_to_markdown(doc, title=title)
        return Response(
            body.encode("utf-8"),
            media_type="text/markdown; charset=utf-8",
            headers={"Content-Disposition": f'attachment; filename="{safe_title}.md"'},
        )

    if format == "txt":
        body = tiptap.doc_to_plaintext(doc, title=title)
        return Response(
            body.encode("utf-8"),
            media_type="text/plain; charset=utf-8",
            headers={"Content-Disposition": f'attachment; filename="{safe_title}.txt"'},
        )

    if format == "json":
        import json

        return Response(
            json.dumps({"title": title, "doc": doc}, ensure_ascii=False, indent=2).encode(),
            media_type="application/json",
            headers={"Content-Disposition": f'attachment; filename="{safe_title}.json"'},
        )

    if format == "html":
        return Response(
            _to_html(doc, title).encode("utf-8"),
            media_type="text/html; charset=utf-8",
            headers={"Content-Disposition": f'attachment; filename="{safe_title}.html"'},
        )

    return _to_docx(doc, title, safe_title)


def _to_html(doc: dict[str, Any], title: str) -> str:
    import html as html_module

    parts = [
        "<!doctype html><html><head><meta charset='utf-8'>",
        f"<title>{html_module.escape(title)}</title>",
        "<style>body{font-family:system-ui,sans-serif;max-width:46rem;margin:3rem auto;"
        "padding:0 1.5rem;line-height:1.65;color:#1a1a1a}"
        "pre{background:#f5f5f5;padding:1rem;overflow-x:auto;border-radius:6px}"
        "table{border-collapse:collapse;width:100%}td,th{border:1px solid #ddd;padding:.5rem}"
        "blockquote{border-left:3px solid #ccc;margin:0;padding-left:1rem;color:#555}"
        "</style></head><body>",
        f"<h1>{html_module.escape(title)}</h1>",
    ]
    for block in tiptap.doc_to_blocks(doc):
        escaped = html_module.escape(block.text)
        if block.type == "heading":
            level = min(6, block.level or 2)
            parts.append(f"<h{level}>{escaped}</h{level}>")
        elif block.type == "list":
            items = "".join(f"<li>{html_module.escape(i)}</li>" for i in block.items if i)
            parts.append(f"<ul>{items}</ul>")
        elif block.type == "code":
            parts.append(f"<pre><code>{escaped}</code></pre>")
        elif block.type == "latex":
            parts.append(f"<p><code>$${escaped}$$</code></p>")
        elif block.type == "mermaid":
            parts.append(f"<pre><code>{escaped}</code></pre>")
        elif block.type == "callout":
            parts.append(
                f"<blockquote><strong>{block.variant or 'Note'}</strong><br>{escaped}</blockquote>"
            )
        elif block.type == "table":
            head = "".join(f"<th>{html_module.escape(h)}</th>" for h in block.headers)
            body = "".join(
                "<tr>" + "".join(f"<td>{html_module.escape(c)}</td>" for c in row) + "</tr>"
                for row in block.rows
            )
            parts.append(f"<table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>")
        else:
            parts.append(f"<p>{escaped}</p>")
    parts.append("</body></html>")
    return "\n".join(parts)


def _to_docx(doc: dict[str, Any], title: str, safe_title: str) -> Response:
    from docx import Document
    from docx.shared import Pt

    document = Document()
    document.add_heading(title, level=0)

    for block in tiptap.doc_to_blocks(doc):
        if block.type == "heading":
            document.add_heading(block.text, level=min(4, block.level or 2))
        elif block.type == "list":
            for item in block.items or block.text.split("\n"):
                if item.strip():
                    document.add_paragraph(item.strip(), style="List Bullet")
        elif block.type in {"code", "latex", "mermaid"}:
            paragraph = document.add_paragraph()
            run = paragraph.add_run(block.text)
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        elif block.type == "callout":
            document.add_paragraph(
                f"{(block.variant or 'note').title()}: {block.text}", style="Intense Quote"
            )
        elif block.type == "table" and (block.headers or block.rows):
            columns = len(block.headers) or max((len(r) for r in block.rows), default=1)
            table = document.add_table(rows=0, cols=columns)
            table.style = "Light Grid Accent 1"
            if block.headers:
                cells = table.add_row().cells
                for index, header in enumerate(block.headers[:columns]):
                    cells[index].text = header
            for row in block.rows:
                cells = table.add_row().cells
                for index in range(columns):
                    cells[index].text = row[index] if index < len(row) else ""
        elif block.text.strip():
            document.add_paragraph(block.text)

    buffer = io.BytesIO()
    document.save(buffer)
    return Response(
        buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{safe_title}.docx"'},
    )
