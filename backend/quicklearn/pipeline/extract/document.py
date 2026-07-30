"""Document extraction: PDF, DOCX, PPTX, plain text.

PDF uses PyMuPDF for the text layer — fast, and crucially no ML dependency, which is what keeps
the base image near 1.5 GB (PRD §6.3). When the text layer is too sparse to be real the page is
a scan, and we rasterise and hand it to a vision model instead of returning a blank note.

DOCX and PPTX preserve heading hierarchy and slide boundaries respectively, because those are
the structural cues the outline pass uses to find section breaks.
"""

from __future__ import annotations

import asyncio
import base64
import logging
from pathlib import Path

from ...errors import PipelineError
from ...providers.base import ImagePart, Message
from .base import ExtractResult, RawUnit

log = logging.getLogger(__name__)

#: Below this many characters per page, assume the text layer is decorative and the real
#: content is an image. Chosen to sit well under a sparse slide but above a page number alone.
SCAN_THRESHOLD_CHARS = 120
#: Rasterisation scale for the OCR fallback. 2x gives ~150 DPI, enough for a vision model
#: without producing images too large to send.
RASTER_SCALE = 2.0
MAX_OCR_PAGES = 40


def _extract_pdf_sync(path: Path) -> tuple[list[RawUnit], int, list[int]]:
    """Text layer per page. Returns units, page count, and pages that look scanned."""
    import fitz

    units: list[RawUnit] = []
    scanned: list[int] = []

    with fitz.open(str(path)) as doc:
        page_count = doc.page_count
        for page_index in range(page_count):
            page = doc.load_page(page_index)
            page_no = page_index + 1

            # Block-level extraction keeps paragraph boundaries and gives a bbox per block,
            # which is what document click-to-source anchors to.
            blocks = page.get_text("blocks")
            page_chars = 0
            for block in sorted(blocks, key=lambda b: (round(b[1], 1), round(b[0], 1))):
                text = str(block[4]).strip() if len(block) > 4 else ""
                if not text:
                    continue
                page_chars += len(text)
                units.append(
                    RawUnit(
                        text=text,
                        page_no=page_no,
                        bbox=[float(block[0]), float(block[1]), float(block[2]), float(block[3])],
                    )
                )

            if page_chars < SCAN_THRESHOLD_CHARS:
                scanned.append(page_no)

    return units, page_count, scanned


def _rasterise_pages_sync(path: Path, pages: list[int]) -> list[tuple[int, bytes]]:
    import fitz

    images: list[tuple[int, bytes]] = []
    with fitz.open(str(path)) as doc:
        for page_no in pages:
            page = doc.load_page(page_no - 1)
            pixmap = page.get_pixmap(matrix=fitz.Matrix(RASTER_SCALE, RASTER_SCALE))
            images.append((page_no, pixmap.tobytes("png")))
    return images


async def extract_pdf(
    path: Path, *, source_id: int | None = None, job_id: int | None = None
) -> ExtractResult:
    units, page_count, scanned = await asyncio.to_thread(_extract_pdf_sync, path)

    # Only treat it as a scan if most of the document is empty. A single image-only page in an
    # otherwise text-bearing PDF is a figure, not a scan, and OCRing it is wasted money.
    if scanned and page_count and len(scanned) > page_count * 0.5:
        log.info(
            "%s looks scanned (%d/%d sparse pages); using the vision fallback",
            path.name,
            len(scanned),
            page_count,
        )
        ocr_units = await _ocr_pages(path, scanned[:MAX_OCR_PAGES], source_id, job_id)
        if ocr_units:
            # Keep whatever real text layer existed, and add the OCR output.
            units = [u for u in units if u.page_no not in set(scanned)] + ocr_units
            units.sort(key=lambda u: u.page_no or 0)

    if not units:
        raise PipelineError(
            f"No readable text found in '{path.name}'. If it is a scan, configure a "
            "vision-capable model (OpenAI, Anthropic or Gemini) in Settings → Model routing."
        )

    return ExtractResult(units=units, page_count=page_count, provider="pymupdf")


async def _ocr_pages(
    path: Path, pages: list[int], source_id: int | None, job_id: int | None
) -> list[RawUnit]:
    """Rasterise and transcribe pages with a vision model."""
    from ... import prompts
    from ...providers import gateway, registry

    try:
        resolved = registry.get_llm("standard")
    except Exception as exc:
        log.warning("no LLM configured for the scanned-PDF fallback: %s", exc)
        return []

    if not resolved.provider.supports_vision:
        log.warning("model %s cannot read images; scanned pages will be skipped", resolved.model)
        return []

    images = await asyncio.to_thread(_rasterise_pages_sync, path, pages)
    instruction = prompts.render("image_describe", hint=f"Page from {path.name}")
    semaphore = asyncio.Semaphore(3)

    async def run(page_no: int, data: bytes) -> RawUnit | None:
        image_part: ImagePart = {
            "type": "image",
            "media_type": "image/png",
            "data": base64.b64encode(data).decode(),
        }
        message: Message = {
            "role": "user",
            "content": [{"type": "text", "text": instruction}, image_part],
        }
        async with semaphore:
            try:
                text = await gateway.complete_text(
                    "ocr_page", [message], source_id=source_id, job_id=job_id, max_tokens=2048
                )
            except Exception as exc:  # a failed page should not fail the whole document
                log.warning("OCR failed on page %d: %s", page_no, exc)
                return None
        return RawUnit(text=text.strip(), page_no=page_no) if text.strip() else None

    results = await asyncio.gather(*(run(p, d) for p, d in images))
    return [unit for unit in results if unit is not None]


def _extract_docx_sync(path: Path) -> list[RawUnit]:
    from docx import Document

    document = Document(str(path))
    units: list[RawUnit] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = (paragraph.style.name or "").lower() if paragraph.style else ""
        # Markdown-mark headings so the outline pass can see the author's own structure
        # instead of having to infer it from prose.
        if style.startswith("heading"):
            level = "".join(c for c in style if c.isdigit()) or "1"
            text = f"{'#' * min(6, int(level))} {text}"
        units.append(RawUnit(text=text))

    for table in document.tables:
        rows = [
            " | ".join(cell.text.strip() for cell in row.cells)
            for row in table.rows
            if any(cell.text.strip() for cell in row.cells)
        ]
        if rows:
            units.append(RawUnit(text="\n".join(rows)))
    return units


async def extract_docx(path: Path) -> ExtractResult:
    units = await asyncio.to_thread(_extract_docx_sync, path)
    if not units:
        raise PipelineError(f"'{path.name}' contains no readable text.")
    return ExtractResult(units=units, provider="python-docx")


def _extract_pptx_sync(path: Path) -> tuple[list[RawUnit], int]:
    from pptx import Presentation

    presentation = Presentation(str(path))
    units: list[RawUnit] = []
    slide_count = 0

    for index, slide in enumerate(presentation.slides, start=1):
        slide_count = index
        pieces: list[str] = []
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            text = shape.text_frame.text.strip()
            if text:
                pieces.append(text)

        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()

        if not pieces and not notes:
            continue

        body = "\n".join(pieces)
        if notes:
            # Speaker notes usually carry the actual explanation, so they are content, not
            # metadata — but they are labelled so the model knows they were not on screen.
            body = f"{body}\n\n[Speaker notes] {notes}".strip()

        # page_no carries the slide number, which is what document anchoring displays.
        units.append(RawUnit(text=f"## Slide {index}\n{body}", page_no=index))

    return units, slide_count


async def extract_pptx(path: Path) -> ExtractResult:
    units, slides = await asyncio.to_thread(_extract_pptx_sync, path)
    if not units:
        raise PipelineError(f"'{path.name}' contains no readable text.")
    return ExtractResult(units=units, page_count=slides, provider="python-pptx")


def extract_text(content: str) -> ExtractResult:
    """Pasted or plain-text input — no decoding needed, split on blank lines."""
    blocks = [block.strip() for block in content.split("\n\n")]
    units = [RawUnit(text=block) for block in blocks if block]
    if not units and content.strip():
        units = [RawUnit(text=content.strip())]
    return ExtractResult(units=units, provider="text")
